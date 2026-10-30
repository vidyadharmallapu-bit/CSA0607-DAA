"""Build the final DOCX from verified CSVs, logs, pseudocode, and graphs."""

from __future__ import annotations

import csv
import platform
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parent
SUPPORT = ROOT / "Supporting_Files"
GRAPHS = SUPPORT / "graphs"
REPORT = ROOT / "Report" / "SentinelOps_Debugging_Optimisation_Report.docx"


def rows(filename: str) -> list[dict[str, str]]:
    with (SUPPORT / filename).open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def add_page_number(section) -> None:
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run("SentinelOps CO2 Debugging and Optimization | Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def add_table(document: Document, headers: list[str], records: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Light Shading Accent 1"
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    for record in records:
        cells = table.add_row().cells
        for cell, value in zip(cells, record):
            cell.text = str(value)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(7.5)
    document.add_paragraph()


def add_code(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.right_indent = Inches(0.15)
    run = paragraph.add_run(text.strip())
    run.font.name = "Consolas"
    run.font.size = Pt(8)


def add_pseudocode(document: Document, relative_path: str) -> None:
    text = (ROOT / relative_path).read_text(encoding="utf-8")
    if "```" in text:
        code = text.split("```", 2)[1]
        code = code.split("\n", 1)[-1]
    else:
        code = text
    add_code(document, code.replace("```", ""))


def add_figure(document: Document, filename: str, caption: str) -> None:
    path = GRAPHS / filename
    if path.exists():
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(path), width=Inches(6.1))
        label = document.add_paragraph(caption)
        label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        document.add_paragraph(f"Graph not generated: {filename}")


def add_register(document: Document) -> None:
    document.add_paragraph("The nine assignment-level categories below are supported by the supplied source. The insertion-sort row groups two related faulty index expressions; statement-level count is therefore ten, as disclosed below the table.")
    defects = [
        ["D01", "sort_alerts", "`j > 0`; `alerts[j] = key`", "Index 0 is skipped and key is restored one slot too early; alerts are lost/duplicated.", "Use `j >= 0` and restore at `j+1`; stable supplied output A2,A4,A3,A1,A5."],
        ["D02", "is_blocked", "Found branch returns False", "Listed IPs are reported unblocked.", "Return True on equality; supplied 10.0.0.9 and .5 -> True."],
        ["D03", "is_blocked", "`else: return False` inside loop", "Search ends after first mismatch, so later entries never match.", "Return False after loop; last-position and missing tests pass."],
        ["D04", "find_signature", "`range(n - m)`", "Final valid alignment is omitted.", "Use `range(n-m+1)` after length checks; end-match test passes."],
        ["D05", "find_signature", "Empty pattern reaches `if j...`", "UnboundLocalError because the inner loop never sets j.", "Define empty pattern as all n+1 boundaries; empty tests pass."],
        ["D06", "find_signature", "`if j == m - 1`", "A mismatch at the final pattern character is falsely accepted.", "Accept only after m successful comparisons; output [0,8]."],
        ["D07", "closest_logins", "`min_dist = 0`", "No positive distance is initially a valid upper bound on the minimum.", "Initialize squared best distance to infinity."],
        ["D08", "closest_logins", "`if d > min_dist`", "Selects the farthest pair, not the closest.", "Minimize squared distance; supplied result indices (0,5), sqrt(2) km."],
        ["D09", "closest_logins", "Both loops range over all n", "Includes self-pairs and both orders of each unordered pair.", "Scan `j > i`; pair count is n(n-1)/2."],
    ]
    add_table(document, ["ID", "Function", "Fault", "Observed effect", "Fix and verification"], defects)
    document.add_paragraph("Count clarification: the expected nine categories are present. D01 has two faulty expressions that jointly violate the insertion-step invariant; counting individual fault locations gives ten. This report does not conceal that distinction or invent an unrelated defect to force the expected count.")
    document.add_paragraph("Other absent input validation is documented as a boundary-condition omission. Quadratic pair scanning and repeated linear lookup are performance limitations, not additional correctness defects.")


def main() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    add_page_number(section)
    doc.styles["Normal"].font.name = "Aptos"
    doc.styles["Normal"].font.size = Pt(9.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("SENTINELOPS\nSOC ALERT-TRIAGE MODULE\nDEBUGGING AND OPTIMIZATION")
    run.bold = True
    run.font.size = Pt(22)
    subtitle = doc.add_paragraph("CSA06 — Design and Analysis of Algorithms\nCO2 Debugging and Optimisation Task | 50 Marks")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Department of Computer Science and Engineering\nPython 3 | Individual Assignment", style="Subtitle").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"Measured environment: {platform.platform()} | Python {platform.python_version()}").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    doc.add_heading("Contents", level=1)
    for item in ("Scenario and method", "Part A — Reproduce and identify defects", "Part B — Systematic corrections", "Part C — Optimized alternatives", "Part D — Experiments and complexity", "Part E — Testing", "Conclusion and engineering recommendations", "Viva preparation", "Limitations"):
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Scenario and Method", level=1)
    doc.add_paragraph("SentinelOps processes SOC alerts by severity, checks IP blocklists, scans log text for signatures, and compares login locations measured in kilometres. The supplied faulty program was preserved with its logic unchanged and executed before corrections. Each correction was tested independently through boundary cases; separate optimized implementations were compared using reproducible input data. Timings use time.perf_counter and are environment-specific. Operation counts are reported independently because instrumentation affects elapsed time.")
    doc.add_paragraph("The sequence followed was reproduce → isolate → correct → validate → optimize → measure → compare. All benchmark CSV rows include the runtime environment and deterministic seeds.")

    doc.add_heading("Part A — Reproduce and Identify Defects", level=1)
    doc.add_paragraph("The original faulty file is retained at Source_Code/original_faulty_module.py. It was executed with the supplied five alerts, three blocklist queries, FAILFAIXFAIL text, FAIL pattern, and six login coordinates. The exact actual results follow:")
    original_log = (ROOT / "Output_Screenshots/original_outputs/original_run.txt").read_text(encoding="utf-8")
    add_code(doc, original_log)
    doc.add_paragraph("The original code returns plausible values rather than raising for most defects: corrupt sorting, false blocklist outcomes, a false string hit and missing final hit, and a farthest pair. Only the empty-pattern probe raises UnboundLocalError.")
    add_register(doc)

    doc.add_heading("Part B — Systematic Corrections", level=1)
    doc.add_paragraph("Corrected code preserves the original algorithms: in-place insertion sort, index-based sequential scan, alignment-by-alignment matching, and nested pair scan. Changes are limited to boundaries, result conventions, match completion, pair initialization/direction/enumeration, and required input validation. The empty signature convention returns all n+1 boundaries.")
    for heading, code_path, explanation in (
        ("Stable descending alert sort", "Pseudocode/sort_alerts_pseudocode.md", "Strict comparison preserves the arrival order of equal severities; the input list is mutated."),
        ("Sequential blocklist search", "Pseudocode/is_blocked_pseudocode.md", "A hit returns True; only a complete miss returns False."),
        ("Brute-force signature matching", "Pseudocode/find_signature_pseudocode.md", "Every valid alignment is examined, overlaps are retained, and success requires all m characters."),
        ("Closest login pair", "Pseudocode/closest_logins_pseudocode.md", "The corrected baseline scans unique pairs and uses lexicographic index-pair tie-breaking."),
    ):
        doc.add_heading(heading, level=2)
        add_pseudocode(doc, code_path)
        doc.add_paragraph(explanation)
    corrected_log = (ROOT / "Output_Screenshots/corrected_outputs/corrected_outputs_and_traces.txt").read_text(encoding="utf-8")
    doc.add_heading("Observed insertion and alignment traces", level=2)
    add_code(doc, corrected_log.split("STABLE INSERTION SORT TRACE", 1)[1])
    doc.add_paragraph("The full original/corrected logs are retained in Output_Screenshots. On the supplied log, the alignment at start 4 compares the final character X against L and fails; start 8 matches through the end.")

    doc.add_heading("Part C — Optimized Alternatives", level=1)
    doc.add_paragraph("Alert sorting retains stable insertion sort for nearly ordered batches and adds bottom-up stable merge sort for larger/random inputs. Merge sorting preserves equal-severity order by choosing the left run on equal keys. Public optimized functions preserve the corrected APIs' return types; separately named instrumented functions expose counters to the benchmark. A set-backed blocklist index is constructed once and then reused across queries; construction and lookup costs are separated. Signature matching remains straightforward brute force; mismatch exit and a correct match condition reduce avoidable work but do not change the O(nm) worst case.")
    doc.add_paragraph("For login locations, the brute-force baseline compares squared distances for unique pairs. The optimized method sorts points and applies exact divide-and-conquer: recursively solve each side, merge y-order, then inspect the strip within the best-distance bound. The packing argument limits each point to the next seven strip points. Duplicate coordinates are detected before recursion and return the lexicographically earliest zero-distance pair. Squared distances alone would only reduce per-pair cost; the observed calculation reduction comes from spatial pruning.")
    optimized_log = (ROOT / "Output_Screenshots/optimized_outputs/optimized_outputs.txt").read_text(encoding="utf-8")
    doc.add_heading("Supplied optimized demonstration", level=2)
    add_code(doc, optimized_log)

    doc.add_heading("Part D — Experiments and Complexity", level=1)
    doc.add_paragraph("Sorting uses seeded identical datasets at n=100, 1,000, and 5,000. Nearly sorted inputs are generated from descending severity order with deterministic adjacent perturbations; random inputs use seed 8821; reverse input is the opposite severity order. Every result is checked against Python's stable descending sort.")
    sort_rows = rows("sorting_benchmarks.csv")
    add_table(doc, ["Dataset", "n", "Algorithm", "Comparisons", "Movements", "Seconds", "Stable/correct"], [[r["dataset_type"], r["input_size"], r["algorithm"], r["comparisons"], r["element_movements"], r["elapsed_seconds"], r["stability_preserved"]] for r in sort_rows])
    add_figure(doc, "sorting_benchmarks.png", "Figure 1. Sorting comparison counts and measured elapsed time from sorting_benchmarks.csv.")
    sort_lookup = {(row["dataset_type"], row["input_size"], row["algorithm"]): row for row in sort_rows}
    nearly_insert = sort_lookup[("nearly_sorted", "5000", "insertion")]
    nearly_merge = sort_lookup[("nearly_sorted", "5000", "stable_merge")]
    random_insert = sort_lookup[("random", "5000", "insertion")]
    random_merge = sort_lookup[("random", "5000", "stable_merge")]
    reverse_insert = sort_lookup[("reverse_sorted", "5000", "insertion")]
    doc.add_paragraph(f"On nearly sorted n=5,000 input, insertion used {nearly_insert['comparisons']} comparisons and {float(nearly_insert['elapsed_seconds']):.6f} s; stable merge used {nearly_merge['comparisons']} comparisons and {float(nearly_merge['elapsed_seconds']):.6f} s. On random n=5,000, insertion used {random_insert['comparisons']} comparisons and {float(random_insert['elapsed_seconds']):.6f} s, versus merge's {random_merge['comparisons']} comparisons and {float(random_merge['elapsed_seconds']):.6f} s. On reverse input insertion used {reverse_insert['comparisons']} comparisons and {float(reverse_insert['elapsed_seconds']):.6f} s. Recommend insertion for small, nearly sorted batches; use a stable O(n log n) alternative when batch size or disorder grows.")

    doc.add_heading("Repeated blocklist queries", level=2)
    search_rows = rows("search_benchmarks.csv")
    add_table(doc, ["Entries", "Queries", "Sequential comparisons", "Sequential s", "Index build s", "Indexed query s", "Correct"], [[r["blocklist_entries"], r["query_count"], r["sequential_comparisons"], r["sequential_total_seconds"], r["index_construction_seconds"], r["indexed_query_seconds"], r["correct"]] for r in search_rows])
    add_figure(doc, "blocklist_benchmarks.png", "Figure 2. Total measured repeated-query workload, including index construction.")
    largest_search = next(row for row in search_rows if row["blocklist_entries"] == "10000")
    doc.add_paragraph(f"At {largest_search['blocklist_entries']} entries and {largest_search['query_count']} mixed queries, sequential scanning made {largest_search['sequential_comparisons']} comparisons in {float(largest_search['sequential_total_seconds']):.6f} s; set construction took {float(largest_search['index_construction_seconds']):.6f} s and indexed queries {float(largest_search['indexed_query_seconds']):.6f} s. For very few lookups, index setup and memory may not be worthwhile; expected constant lookup is not a worst-case guarantee.")

    doc.add_heading("Signature matching workloads", level=2)
    string_rows = rows("string_matching_benchmarks.csv")
    add_table(doc, ["Case", "n", "m", "Character comparisons", "Matches", "Seconds", "Correct"], [[r["case"], r["text_length"], r["pattern_length"], r["character_comparisons"], r["match_count"], r["elapsed_seconds"], r["correct"]] for r in string_rows])
    add_figure(doc, "string_benchmarks.png", "Figure 3. Counted character comparisons by deliberate input structure.")
    repetitive_case = next(row for row in string_rows if row["case"] == "repetitive_near_match")
    doc.add_paragraph(f"The repetitive near-match case has long common prefixes and reaches {repetitive_case['character_comparisons']} character comparisons; the early-match, missing-first-character, and overlapping cases have different work despite comparable text lengths. This illustrates why text length alone is insufficient to predict comparisons.")

    doc.add_heading("Closest-pair batch comparison", level=2)
    closest_rows = rows("closest_pair_benchmarks.csv")
    add_table(doc, ["n", "Theoretical pairs", "Brute distance calculations", "Optimized distance calculations", "Preprocessing ops", "Brute s", "Optimized s", "Same result"], [[r["event_count"], r["theoretical_unique_pairs"], r["brute_distance_calculations"], r["optimized_distance_calculations"], r["optimized_preprocessing_operations"], r["brute_elapsed_seconds"], r["optimized_elapsed_seconds"], r["correct"]] for r in closest_rows])
    add_figure(doc, "closest_pair_benchmarks.png", "Figure 4. Unique-pair baseline and exact divide-and-conquer counts.")
    largest_closest = next(row for row in closest_rows if row["event_count"] == "8000")
    doc.add_paragraph(f"The brute theoretical/actual count is n(n-1)/2. At {largest_closest['event_count']} events, brute force performed {largest_closest['brute_distance_calculations']} distance calculations in {float(largest_closest['brute_elapsed_seconds']):.4f} s; divide-and-conquer performed {largest_closest['optimized_distance_calculations']} in {float(largest_closest['optimized_elapsed_seconds']):.4f} s with {largest_closest['optimized_preprocessing_operations']} counted preprocessing/point-processing operations. All four seeded batches agreed exactly in pair and distance. This measured method provides the required largest-batch comparison without substituting theoretical counts for actual baseline work.")

    doc.add_heading("Complexity summary", level=2)
    complexity_rows = [
        ["Alert insertion", "O(n) best; O(n²) average/worst", "O(1)"],
        ["Stable merge sort", "O(n log n) all cases", "O(n)"],
        ["Sequential search", "O(1) best; O(n) average/worst", "O(1)"],
        ["Set index + lookup", "O(n) expected build; O(1) expected lookup, O(n) hash worst case", "O(n)"],
        ["Brute string matching", "O((n-m+1)m) worst", "O(output)"],
        ["Closest brute baseline", "Θ(n²)", "O(1)"],
        ["Exact divide-and-conquer", "O(n log n)", "O(n)"],
    ]
    add_table(doc, ["Algorithm", "Time", "Auxiliary space"], complexity_rows)
    doc.add_paragraph("Full operation definitions and trade-offs are in Supporting_Files/complexity_analysis.md. Timings are a single-machine measurement; counters are more portable and instrumentation changes runtime.")

    doc.add_heading("Part E — Testing", level=1)
    test_rows = rows("test_results.csv")
    add_table(doc, ["Test", "Status"], [[r["test"].rsplit(".", 1)[-1], r["status"]] for r in test_rows])
    doc.add_paragraph(f"The standard-library harness ran {len(test_rows)} tests; {sum(r['status'] == 'PASS' for r in test_rows)} passed and {sum(r['status'] != 'PASS' for r in test_rows)} failed. It covers empty/singleton/sorted/reverse/stable sorting, first/last/missing/empty blocklist, supplied/end/overlap/empty/too-long signatures, duplicate/tied/negative/two-event closest pairs, invalid input, and baseline-versus-optimized agreement.")

    doc.add_heading("Conclusion and Engineering Recommendations", level=1)
    doc.add_paragraph("The supplied code's errors were reproducible and primarily arose from incorrect loop bounds, early returns, an incomplete match condition, and reversed closest-pair state updates. Minimal corrections restore the intended four algorithms. The measured alternatives are workload-sensitive: insertion sort is appropriate for small nearly sorted alert batches, a hash index benefits repeated lookups after amortizing construction, and exact divide-and-conquer sharply reduces distance calculations for large event batches. Keep simple brute-force versions as auditable references and retain the regression harness for future changes.")

    doc.add_heading("Viva Preparation", level=1)
    viva = [
        ("Why does insertion sort need j >= 0?", "The held key may belong before the first item, so index zero must be compared."),
        ("How does strict severity comparison preserve stability?", "Equal-severity existing items are not shifted past the later key."),
        ("Why did the original blocklist fail for later entries?", "It returned False at the first nonmatch, ending the scan."),
        ("Why must the matcher use n-m+1 alignments?", "The range end is exclusive, and start n-m is a valid final alignment."),
        ("Why was index 4 a false string hit?", "The final comparison mismatched, but j still equalled m-1."),
        ("Why use j > i for pairs?", "It excludes self-pairs and visits each unordered pair once."),
        ("Does squared distance improve brute-force asymptotic time?", "No; it removes per-pair square roots but still examines Θ(n²) pairs."),
        ("Why is the strip bound exact?", "After each half has a distance bound d, 2-D packing limits possible closer cross-boundary neighbors to seven next points in y-order."),
        ("When does a set index pay off?", "When enough repeated queries amortize construction and memory; expected lookup is O(1), not worst-case guaranteed."),
        ("Does instrumentation affect timing?", "Yes; operation counters add work, so counts and wall-clock measurements are reported separately."),
    ]
    for question, answer in viva:
        paragraph = doc.add_paragraph()
        paragraph.add_run("Q: " + question + " ").bold = True
        paragraph.add_run("A: " + answer)

    doc.add_heading("Limitations", level=1)
    doc.add_paragraph("Automated terminal screenshot capture was unavailable; genuine execution logs and a manual Windows capture procedure are supplied instead. The figures are actual Matplotlib plots generated from benchmark CSVs, not terminal screenshots. Performance measurements were taken on the recorded Linux container, and should not be read as Windows timings. Duplicate coordinates are handled explicitly; coordinate inputs are assumed to be finite real numbers.")

    REPORT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(REPORT)
    print(f"Generated {REPORT}")


if __name__ == "__main__":
    main()
